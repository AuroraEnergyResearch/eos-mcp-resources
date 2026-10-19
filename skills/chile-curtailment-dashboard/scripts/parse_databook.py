#!/usr/bin/env python3
"""Extract the dashboard numbers from an Aurora Chile PRMF "Data" databook.

Usage:
  python parse_databook.py <databook.xlsx> <out.json> [--scenario Central] [--start 2027] [--end 2035]

Rows are located by their labels, not fixed row numbers, so the script keeps
working when a new quarterly databook shifts rows around.

Definitions (as agreed with the dashboard owner):
  * Renewable generation = Solar PV + Onshore wind + Hydro + Biomass/other RES,
    from the scenario sheet's "Generation / Generación" block (Chile, TWh).
  * Curtailed TWh per hub = gen x p / (1 - p) for Solar PV and Onshore wind,
    p = "Curtailment - fleet wide" % by hub on the scenario sheet,
    gen = hub generation on the "Zonal <scenario>" sheet.
    The "Scenario Comparison" tab is deliberately NOT used.
  * Share curtailed = curtailed / (solar + wind generation + curtailed).
  * Battery capacity = "Battery storage" GW on the Zonal sheet, summed over hubs.
"""
import argparse, json, sys
import openpyxl

HUBS = ["Tarapacá", "Crucero", "Atacama", "Diego de Almagro", "Cardones", "Maitencillo",
        "Pan de Azucar", "Choapa", "Quillota", "Polpaico", "Alto Jahuel", "Charrúa",
        "Concepción", "Temuco", "Osorno", "Puerto Montt"]
RE_TECHS = ["Solar PV", "Onshore wind", "Hydro", "Biomass/ other RES"]


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def s(v):
    return str(v).strip() if v is not None else ""


def year_cols(row, years):
    """Map year -> column index from a header row (years may be int or str)."""
    m = {}
    for i, v in enumerate(row):
        try:
            y = int(float(v))
        except (TypeError, ValueError):
            continue
        if y in years:
            m[y] = i
    return m


def find_header_after(rows, start, years):
    for i in range(start, min(start + 6, len(rows))):
        m = year_cols(rows[i], years)
        if len(m) == len(years):
            return i, m
    raise ValueError(f"No year header found near row {start + 1}")


def parse(path, scenario, years):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if scenario not in wb.sheetnames or f"Zonal {scenario}" not in wb.sheetnames:
        sys.exit(f"Sheets '{scenario}' / 'Zonal {scenario}' not found. Sheets: {wb.sheetnames}")
    sc = [tuple(r) for r in wb[scenario].iter_rows(values_only=True)]
    zn = [tuple(r) for r in wb[f"Zonal {scenario}"].iter_rows(values_only=True)]

    # --- national generation block -------------------------------------------
    gen = {}
    for i, r in enumerate(sc):
        if any(s(c).startswith("Generation / Generación") for c in r[:3]):
            hi, yc = find_header_after(sc, i - 2, years)
            for r2 in sc[hi + 1:hi + 20]:
                if s(r2[2]) and s(r2[5]) == "TWh":
                    gen[s(r2[2])] = [num(r2[yc[y]]) for y in years]
                if s(r2[1]).startswith("Notes"):
                    break
            break
    missing = [t for t in RE_TECHS if t not in gen]
    if missing:
        sys.exit(f"Generation rows not found for {missing}")

    # --- curtailment % by hub -------------------------------------------------
    curt = {"Solar PV": {}, "Onshore wind": {}}
    start = next(i for i, r in enumerate(sc) if s(r[1]).startswith("Curtailment - fleet wide"))
    # the header row with years sits a few rows above the block
    yc = None
    for j in range(start, max(0, start - 150), -1):
        m = year_cols(sc[j], years)
        if len(m) == len(years):
            yc = m
            break
    tech = None
    for r in sc[start:start + 60]:
        if s(r[1]).startswith("Notes") or s(r[1]).startswith("Percentage"):
            break
        if s(r[2]) in curt:
            tech = s(r[2])
        if tech and s(r[4]) and s(r[5]) == "%":
            curt[tech][s(r[4])] = [num(r[yc[y]]) / 100 for y in years]

    # --- zonal capacity and generation ---------------------------------------
    def zonal_block(title):
        i0 = next(i for i, r in enumerate(zn) if s(r[0]).startswith(title))
        hi, zc = find_header_after(zn, i0, years)
        out, hub = {}, None
        for r in zn[hi + 1:]:
            if s(r[0]) and s(r[0]) != title:
                break
            if s(r[1]).startswith("Note"):
                break
            if s(r[1]):
                hub = s(r[1])
            if hub and s(r[2]) and s(r[2]) != "Variable":
                out.setdefault(hub, {})[s(r[2])] = [num(r[zc[y]]) for y in years]
        return out

    cap = zonal_block("Capacity / Capacidad")
    zg = zonal_block("Generation / Generación")
    if set(HUBS) - set(cap):
        sys.exit(f"Hubs missing from zonal sheet: {set(HUBS) - set(cap)}")

    n = len(years)
    hubs = []
    warnings = []
    for h in HUBS:
        rec = {"name": h, "bat": [], "curt": [], "pot": []}
        for k in range(n):
            tot_c = tot_g = 0.0
            for t in ("Solar PV", "Onshore wind"):
                g = zg[h].get(t, [0.0] * n)[k]
                p = curt[t].get(h, [0.0] * n)[k]
                if g > 0 and h not in curt[t]:
                    warnings.append(f"{h} has {t} generation but no curtailment row (treated as 0%)")
                c = g * p / (1 - p) if p < 1 else 0.0
                tot_c += c
                tot_g += g
            rec["curt"].append(round(tot_c, 4))
            rec["pot"].append(round(tot_g + tot_c, 4))
            rec["bat"].append(round(cap[h].get("Battery storage", [0.0] * n)[k], 3))
        hubs.append(rec)

    nat = {
        "re": [round(sum(gen[t][k] for t in RE_TECHS), 2) for k in range(n)],
        "curt": [round(sum(h["curt"][k] for h in hubs), 3) for k in range(n)],
        "pot": [round(sum(h["pot"][k] for h in hubs), 3) for k in range(n)],
        "bat": [round(sum(h["bat"][k] for h in hubs), 2) for k in range(n)],
    }
    nat["share"] = [round(100 * c / p, 2) if p else 0 for c, p in zip(nat["curt"], nat["pot"])]
    return {"years": years, "hubs": hubs, "national": nat}, sorted(set(warnings))


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("out")
    ap.add_argument("--scenario", default="Central")
    ap.add_argument("--start", type=int, default=2027)
    ap.add_argument("--end", type=int, default=2035)
    a = ap.parse_args()
    data, warns = parse(a.xlsx, a.scenario, list(range(a.start, a.end + 1)))
    json.dump(data, open(a.out, "w", encoding="utf-8"), ensure_ascii=False)
    nat = data["national"]
    print("Year      RE TWh  Curt TWh  Share %  Battery GW")
    for k, y in enumerate(data["years"]):
        print(f"{y}  {nat['re'][k]:8.1f}  {nat['curt'][k]:8.2f}  {nat['share'][k]:7.1f}  {nat['bat'][k]:9.1f}")
    for w in warns:
        print("WARNING:", w)
